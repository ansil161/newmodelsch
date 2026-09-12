"""Real documents, built in memory, for the ingestion tests.

A parser test against a hand-described structure proves nothing about files.
These are genuine PDFs and Word documents — the PDF written byte by byte with
a correct cross-reference table, the DOCX by python-docx itself — so what is
tested is what an administrator's upload would actually contain.
"""

from __future__ import annotations

import io

from docx import Document as new_document
from pypdf import PdfReader, PdfWriter


def make_pdf(pages: list[list[str]]) -> bytes:
    """A PDF with one text line per entry, one page per list."""
    objects: dict[int, bytes] = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    }
    kids: list[int] = []
    next_id = 4
    for lines in pages:
        page_id, content_id = next_id, next_id + 1
        next_id += 2
        operations = ["BT", "/F1 12 Tf", "72 720 Td"]
        for position, line in enumerate(lines):
            if position:
                operations.append("0 -18 Td")
            escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            operations.append(f"({escaped}) Tj")
        operations.append("ET")
        stream = "\n".join(operations).encode("latin-1")
        objects[content_id] = (
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"
        )
        objects[page_id] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>"
        ).encode()
        kids.append(page_id)
    references = " ".join(f"{kid} 0 R" for kid in kids)
    objects[2] = f"<< /Type /Pages /Kids [{references}] /Count {len(kids)} >>".encode()

    output = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for object_id in sorted(objects):
        offsets[object_id] = len(output)
        output += f"{object_id} 0 obj\n".encode() + objects[object_id] + b"\nendobj\n"
    xref = len(output)
    size = max(objects) + 1
    output += f"xref\n0 {size}\n".encode() + b"0000000000 65535 f \n"
    for object_id in range(1, size):
        output += f"{offsets[object_id]:010d} 00000 n \n".encode()
    output += (
        f"trailer\n<< /Size {size} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    ).encode()
    return bytes(output)


def rewrite_pdf(data: bytes, *, outline: list[tuple[str, int]] = (),  # type: ignore[assignment]
                user_password: str | None = None) -> bytes:
    """Add outline entries (title, 0-based page) and/or encryption to a PDF."""
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(data)))
    for title, page in outline:
        writer.add_outline_item(title, page)
    if user_password is not None:
        writer.encrypt(user_password=user_password, owner_password="owner-secret",
                       algorithm="AES-128")
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def make_docx() -> bytes:
    document = new_document()
    document.core_properties.title = "Staff Handbook"
    document.add_heading("Leave policy", level=1)
    document.add_paragraph("Teachers may take twelve days of casual leave a year.")
    document.add_heading("Applying", level=2)
    document.add_paragraph("Submit the form to the office", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Type"
    table.cell(0, 1).text = "Days"
    table.cell(1, 0).text = "Casual"
    table.cell(1, 1).text = "12"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
